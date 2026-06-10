"""
Ingesta de PDFs → corpus de fragmentos para el generador de preguntas.

NOVEDADES de esta versión:
─────────────────────────
- Cada fragmento se etiqueta con un campo `area` (opcional). Esto permite
  filtrar el corpus por subárea dentro de una asignatura.
  Ejemplo: dentro de "biochem" tenemos áreas "metabolismo" y "genetica".
- El mapeo PDF → área se define en backend/materials/<subject>/areas.json
  (opcional). Si no existe ese archivo, `area = null` para todo.

EJEMPLO de areas.json para biochem:
─────────────────────────────────
{
  "metabolismo": [
    "Anabolismo_Hidratos",
    "Catabolismo_lipidos",
    "Metabolismo_Oxidativo",
    "TEMA_11b__Metabolismo_lipi",
    "TEMA_12__Metabolismo_del_Nitro",
    "Tema_13___Integracio"
  ],
  "genetica": [
    "Tema_14__Replicacion",
    "Tema_15__Transcripci",
    "Tema_16__Traducci",
    "Tema_17__Regulaci",
    "Tema_18__Introducci"
  ]
}

Cada string es un PATRÓN (substring) que se busca en el nombre del PDF.
El primer área que matche gana.

USO
───
    cd backend
    python scripts/ingest_pdfs.py
"""
import json
import re
import sys
from pathlib import Path
from typing import List, Dict, Any, Optional

try:
    from pypdf import PdfReader
except ImportError:
    print("ERROR: necesitas instalar pypdf → pip install pypdf")
    sys.exit(1)

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


# ─── Configuración ─────────────────────────────────────────────────────────
CHUNK_WORDS = 500
OVERLAP_WORDS = 50
MIN_CHUNK_WORDS = 80

NOISE_PATTERNS = [
    re.compile(r"^Ángela Martín Cortés\.?\s*Área de Bioquímica.*", re.I),
    re.compile(r"^Universidad Rey Juan Carlos$", re.I),
    re.compile(r"^FCS\.\s*\d{4}-\d{2}$", re.I),
    re.compile(r"^\d{1,3}$"),
    re.compile(r"^Página \d+ de \d+$", re.I),
    re.compile(r"^https?://\S+$"),
]


def clean_text(raw: str) -> str:
    lines = raw.split("\n")
    kept = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if any(p.match(stripped) for p in NOISE_PATTERNS):
            continue
        kept.append(stripped)
    text = " ".join(kept)
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"(\w)-\s+(\w)", r"\1\2", text)
    return text.strip()


def chunk_words(words: List[str], size: int, overlap: int) -> List[List[str]]:
    if not words:
        return []
    chunks = []
    i = 0
    n = len(words)
    while i < n:
        chunk = words[i:i + size]
        if len(chunk) >= MIN_CHUNK_WORDS:
            chunks.append(chunk)
        i += size - overlap
    return chunks


def _file_magic(path: Path) -> str:
    with open(path, "rb") as f:
        head = f.read(4)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return "zip"
    return "unknown"


def process_zip_with_txt(path: Path) -> List[Dict[str, Any]]:
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            txt_names = sorted(
                [n for n in z.namelist() if n.lower().endswith(".txt")],
                key=lambda n: (len(n), n)
            )
            if not txt_names:
                print(f"[saltado: ZIP sin .txt OCR]")
                return []
            slides_text = []
            for name in txt_names:
                try:
                    raw = z.read(name).decode("utf-8", errors="ignore")
                except Exception:
                    continue
                cleaned = clean_text(raw)
                if cleaned:
                    slides_text.append(cleaned)
    except Exception as e:
        print(f"[ERROR al leer ZIP: {e}]")
        return []

    all_words = []
    page_breaks = []
    for slide_num, text in enumerate(slides_text, start=1):
        page_breaks.append((len(all_words), slide_num))
        all_words.extend(text.split())

    if len(all_words) < MIN_CHUNK_WORDS:
        print(f"[saltado: solo {len(all_words)} palabras]")
        return []

    word_chunks = chunk_words(all_words, CHUNK_WORDS, OVERLAP_WORDS)
    out = []
    for chunk_idx, words in enumerate(word_chunks):
        start_word_idx = chunk_idx * (CHUNK_WORDS - OVERLAP_WORDS)
        slide = 1
        for break_idx, p in page_breaks:
            if break_idx <= start_word_idx:
                slide = p
            else:
                break
        out.append({"source": path.name, "page_start": slide, "text": " ".join(words)})
    print(f"(zip-ocr) → {len(out)} fragmentos")
    return out


def process_pptx(path: Path) -> List[Dict[str, Any]]:
    if not HAS_PPTX:
        return []
    try:
        pres = Presentation(str(path))
    except Exception:
        return []

    all_words = []
    page_breaks = []
    for slide_num, slide in enumerate(pres.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        slide_text.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        t = cell.text.strip()
                        if t:
                            slide_text.append(t)
        text = clean_text("\n".join(slide_text))
        if not text:
            continue
        page_breaks.append((len(all_words), slide_num))
        all_words.extend(text.split())

    if len(all_words) < MIN_CHUNK_WORDS:
        print(f"[saltado: solo {len(all_words)} palabras]")
        return []

    word_chunks = chunk_words(all_words, CHUNK_WORDS, OVERLAP_WORDS)
    out = []
    for chunk_idx, words in enumerate(word_chunks):
        start_word_idx = chunk_idx * (CHUNK_WORDS - OVERLAP_WORDS)
        slide = 1
        for break_idx, p in page_breaks:
            if break_idx <= start_word_idx:
                slide = p
            else:
                break
        out.append({"source": path.name, "page_start": slide, "text": " ".join(words)})
    print(f"(pptx) → {len(out)} fragmentos")
    return out


def process_pdf(pdf_path: Path) -> List[Dict[str, Any]]:
    print(f"  · {pdf_path.name}", end=" ", flush=True)
    magic = _file_magic(pdf_path)
    if magic == "zip":
        result = process_pptx(pdf_path)
        if result:
            return result
        return process_zip_with_txt(pdf_path)
    if magic != "pdf":
        print(f"[saltado: formato no reconocido ({magic})]")
        return []
    try:
        reader = PdfReader(str(pdf_path))
    except Exception as e:
        print(f"[ERROR al abrir: {e}]")
        return []

    all_words = []
    page_breaks = []
    for page_num, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:
            raw = ""
        text = clean_text(raw)
        if not text:
            continue
        page_breaks.append((len(all_words), page_num))
        all_words.extend(text.split())

    if len(all_words) < MIN_CHUNK_WORDS:
        print(f"[saltado: solo {len(all_words)} palabras]")
        return []

    word_chunks = chunk_words(all_words, CHUNK_WORDS, OVERLAP_WORDS)
    out = []
    for chunk_idx, words in enumerate(word_chunks):
        start_word_idx = chunk_idx * (CHUNK_WORDS - OVERLAP_WORDS)
        page = 1
        for break_idx, p in page_breaks:
            if break_idx <= start_word_idx:
                page = p
            else:
                break
        out.append({"source": pdf_path.name, "page_start": page, "text": " ".join(words)})
    print(f"→ {len(out)} fragmentos")
    return out


# ──────────────────────────────────────────────────────────────────────────────
# NUEVO: mapeo PDF → área
# ──────────────────────────────────────────────────────────────────────────────

def load_area_map(subject_dir: Path) -> Dict[str, List[str]]:
    """Lee areas.json del directorio si existe. Devuelve dict area → [patrones]."""
    path = subject_dir / "areas.json"
    if not path.exists():
        print(f"  ℹ areas.json NO encontrado en {path}; los fragmentos no se etiquetarán por área.")
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            print(f"  [areas.json] formato inesperado, ignoro")
            return {}
        return data
    except Exception as e:
        print(f"  [areas.json] error al leer ({e}), ignoro")
        return {}


def detect_area(filename: str, area_map: Dict[str, List[str]]) -> Optional[str]:
    """Devuelve la primera área cuyos patrones (substring, case-insensitive) matchean."""
    name_low = filename.lower()
    for area, patterns in area_map.items():
        # Saltamos claves de metadata (las que empiezan por "_")
        if area.startswith("_"):
            continue
        # Patterns puede no ser lista si el JSON está mal
        if not isinstance(patterns, list):
            continue
        for pat in patterns:
            if isinstance(pat, str) and pat.lower() in name_low:
                return area
    return None


def ingest_subject(subject_dir: Path) -> List[Dict[str, Any]]:
    # Deduplicamos por path resuelto: en Windows el sistema de archivos es
    # case-insensitive, así que glob("*.pdf") y glob("*.PDF") devuelven los
    # mismos archivos. Sin deduplicar, cada PDF se procesaría dos veces.
    seen = set()
    files = []
    for f in (list(subject_dir.glob("*.pdf")) +
              list(subject_dir.glob("*.PDF")) +
              list(subject_dir.glob("*.pptx")) +
              list(subject_dir.glob("*.PPTX"))):
        key = str(f.resolve()).lower()
        if key not in seen:
            seen.add(key)
            files.append(f)
    files = sorted(files, key=lambda p: p.name.lower())
    if not files:
        print(f"  (no hay PDFs ni PPTX en {subject_dir})")
        return []

    area_map = load_area_map(subject_dir)
    if area_map:
        print(f"  📂 Mapa de áreas cargado: {list(area_map.keys())}")

    all_fragments = []
    for f in files:
        if f.suffix.lower() == ".pptx":
            print(f"  · {f.name}", end=" ", flush=True)
            frags = process_pptx(f)
        else:
            frags = process_pdf(f)

        area = detect_area(f.name, area_map)
        for fr in frags:
            fr["area"] = area
        all_fragments.extend(frags)

    return all_fragments


def main():
    root = Path(__file__).resolve().parent.parent
    materials = root / "materials"

    if not materials.exists():
        print(f"ERROR: no existe {materials}")
        sys.exit(1)

    subjects = [d for d in materials.iterdir() if d.is_dir()]
    if not subjects:
        print(f"ERROR: no hay subcarpetas en {materials}")
        sys.exit(1)

    for subject_dir in sorted(subjects):
        subject = subject_dir.name
        print(f"\n📚 Procesando asignatura: {subject}")
        fragments = ingest_subject(subject_dir)
        if not fragments:
            print(f"  ⚠ Sin fragmentos para {subject}, saltando.")
            continue

        out_path = materials / f"{subject}_corpus.json"
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump({
                "subject": subject,
                "fragment_count": len(fragments),
                "fragments": fragments,
            }, f, ensure_ascii=False, indent=2)

        total_words = sum(len(fr["text"].split()) for fr in fragments)

        # Resumen por área
        from collections import Counter
        area_counts = Counter(fr.get("area") for fr in fragments)
        area_summary = ", ".join(
            f"{area or 'sin área'}: {n}"
            for area, n in area_counts.most_common()
        )

        print(f"  ✓ {len(fragments)} fragmentos · {total_words:,} palabras → {out_path.name}")
        print(f"     {area_summary}")

    print("\n✓ Ingesta completada.")


if __name__ == "__main__":
    main()